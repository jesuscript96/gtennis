"""Excel limpio: una fila por dato, filtrable, pensado también para cargarlo en la futura app."""
import datetime as dt
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from common import *

OUT = "/Users/jvch/Desktop/GTennis/docs/GTenis - Calendario de torneos 2026 (datos limpios).xlsx"
wb = Workbook()
GREEN = "1F6F4A"
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor=GREEN)
WRAP = Alignment(wrap_text=True, vertical="top")
TITLE = Font(bold=True, size=14, color=GREEN)


def sheet(title, headers, rows, widths, first=False, note=None):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True, color="5D6B64")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.row_dimensions[1].height = 30
        ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="top")
        r0 = 2
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill, c.alignment = HEAD, FILL, Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.alignment = WRAP
            if isinstance(v, dt.date):
                c.number_format = "DD/MM/YYYY"
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    if rows:
        ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
        tb = Table(displayName=re.sub(r"\W", "", title)[:30] or "T", ref=ref)
        tb.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tb)
    return ws


D = dt.date.fromisoformat
conv = sorted([t for t in T if t["kind"] == "convocatoria"], key=lambda t: (t["week"], catkey(t["cat"]), t["name"]))


def cierre_txt(t):
    out = []
    for k, refs in t["cierres"].items():
        out.append(fdate(k))
    return " / ".join(out)


def pl_txt(t):
    xs = []
    for p in t["players"]:
        s = p["nombre"]
        if p["certeza"].startswith("deducido"):
            s += " [deducido]"
        if "pendiente" in p["certeza"]:
            s += " [?]"
        if p["certeza"].startswith("sin identificar") and not p["nombre"].startswith("Sin identificar"):
            s += " [sin identificar]"
        xs.append(s)
    return ", ".join(xs)


def ref_txt(t):
    by = defaultdict(list)
    for p in t["players"]:
        for c in p["entrenador_ref"]:
            by[c].append(p["nombre"])
    return "; ".join(f"{c}: {', '.join(v)}" for c, v in by.items())


# 0. Cómo leer
ws = wb.active
ws.title = "Cómo leer"
lines = [
    ("G Tenis · Calendario de torneos 2026 — datos limpios", TITLE),
    ("Hecho a partir de «CALENDARIO 2026 (2).xlsx»: la hoja maestra y las 12 hojas de grupo, cruzadas entre sí.", None),
    ("", None),
    ("Pestañas", Font(bold=True)),
    ("Semana a semana — una fila por torneo y semana en la que hay al menos un jugador (o un coach) apuntado.", None),
    ("Jugador × torneo — una fila por jugador, torneo y semana. Es la tabla para filtrar por jugador o para cargar en la app.", None),
    ("Por jugador — resumen: cuántas semanas de torneo tiene cada jugador y cuáles.", None),
    ("Coaches — cada viaje con coach acompañante, y los que van solos.", None),
    ("Calendario completo — todos los torneos que aparecen en el Excel, también los que nadie tiene asignados.", None),
    ("Otras anotaciones — entrenos, lesiones, equipos, exámenes, vacaciones y notas de viaje escritas en las hojas.", None),
    ("Incidencias — todo lo que no cuadra y conviene revisar (cierres, nombres dudosos, diferencias entre hojas).", None),
    ("Criterios nombres — cómo se ha decidido quién es quién cuando el Excel pone solo un nombre de pila.", None),
    ("Restos 2025 — filas de 2025 que siguen en las hojas de Salva y Álvaro (no forman parte del calendario 2026).", None),
    ("", None),
    ("Convenciones", Font(bold=True)),
    ("[deducido] — el Excel pone un nombre ambiguo (p. ej. «ERIC», «DIEGO», «CARLA») y se ha asignado a la persona más probable según la hoja. Ver «Criterios nombres».", None),
    ("[?] — en el Excel aparece con interrogación: pendiente de confirmar.", None),
    ("Coach que acompaña — sale de la columna COACH de las hojas-lista (maestra, Pablo, Santi, Salva, Álvaro).", None),
    ("Entrenador de referencia — en las hojas-rejilla (ATP/Challenger/Futures, Futures, Jorge, Mario, Salva 2, Nacho Calvo) cada columna de jugador lleva el nombre de su entrenador; no significa que viaje con él.", None),
    ("Fuente — hoja y celda del Excel original (p. ej. «Maestra C51» = hoja DISTRIBUCION CALENDARIO TODO, celda C51).", None),
    ("Semana — de lunes a domingo. Un torneo que empieza en sábado puede aparecer en semanas distintas según la hoja; se señala en Incidencias.", None),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(i, 1, t)
    if f:
        c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 130

# 1. Semana a semana
rows = []
for t in conv:
    w = D(t["week"])
    rows.append([weekno(t["week"]), w, wrange(t["week"]), pretty(t["name"]), t["cat"], cierre_txt(t),
                 ", ".join(t["acompanan"]), ref_txt(t), len([p for p in t["players"]]), pl_txt(t),
                 " · ".join(t["notas"]), ", ".join(t["sources"]), " | ".join(t["variants"])])
sheet("Semana a semana", ["Nº sem.", "Lunes", "Semana", "Torneo", "Circuito", "Cierre inscripción",
                          "Coach que acompaña", "Entrenador de referencia (hojas-rejilla)", "Nº jug.", "Jugadores",
                          "Notas del Excel", "Fuente (hoja y celda)", "También escrito como"],
      rows, [7, 11, 22, 44, 20, 16, 24, 34, 7, 70, 24, 40, 50])

# 2. Jugador × torneo
rows = []
for t in conv:
    for p in t["players"]:
        rows.append([D(t["week"]), wrange(t["week"]), p["nombre"], pretty(t["name"]), t["cat"],
                     ", ".join(t["acompanan"]), ", ".join(p["entrenador_ref"]), p["certeza"], p["detalle"],
                     cierre_txt(t), ", ".join(p["fuentes"])])
rows.sort(key=lambda r: (r[2], r[0]))
sheet("Jugador × torneo", ["Lunes", "Semana", "Jugador", "Torneo", "Circuito", "Coach que acompaña",
                           "Entrenador de referencia", "Certeza", "Detalle", "Cierre inscripción", "Fuente"],
      rows, [11, 22, 26, 44, 20, 24, 22, 22, 30, 16, 36])

# 3. Por jugador
per = defaultdict(list)
for t in conv:
    for p in t["players"]:
        per[p["nombre"]].append(t)
rows = []
for nm, ts in sorted(per.items(), key=lambda x: (-len({t["week"] for t in x[1]}), x[0])):
    weeks = sorted({t["week"] for t in ts})
    rows.append([nm, len(weeks), len(ts), wshort(weeks[0]), wshort(weeks[-1]),
                 "\n".join(f"{wshort(t['week'])}: {short(t['name'], 60)}" for t in sorted(ts, key=lambda t: t["week"]))])
sheet("Por jugador", ["Jugador", "Semanas con torneo", "Torneos (entradas)", "Primera", "Última", "Detalle"],
      rows, [28, 12, 12, 10, 10, 90])

# 4. Coaches
rows = []
for t in conv:
    for c in t["acompanan"] or []:
        rows.append([c, D(t["week"]), wrange(t["week"]), pretty(t["name"]), t["cat"],
                     ", ".join(p["nombre"] for p in t["players"]), ", ".join(t["sources"])])
rows.sort(key=lambda r: (r[0], r[1]))
sheet("Coaches", ["Coach que acompaña", "Lunes", "Semana", "Torneo", "Circuito", "Jugadores", "Fuente"], rows,
      [26, 11, 22, 44, 20, 70, 36],
      note="Incluye «Sin acompañante» (va solo/sola/solas), «Sí, con coach (sin nombre)» (columna COACH = SI en la hoja de Santi) y «No va coach».")

# 5. Calendario completo
rows = []
for t in sorted(T, key=lambda t: (t["week"], catkey(t["cat"]), t["name"])):
    rows.append([weekno(t["week"]), D(t["week"]), wrange(t["week"]), pretty(t["name"]), t["cat"],
                 "Con jugadores" if t["kind"] == "convocatoria" else "Solo en calendario",
                 len(t["players"]), cierre_txt(t), ", ".join(t["sheets"]), ", ".join(t["sources"]),
                 " | ".join(t["variants"])])
sheet("Calendario completo", ["Nº sem.", "Lunes", "Semana", "Torneo", "Circuito", "Estado", "Nº jug.",
                              "Cierre inscripción", "Hojas", "Fuente", "También escrito como"], rows,
      [7, 11, 22, 46, 20, 16, 7, 16, 30, 40, 50])

# 6. Otras anotaciones
rows = [[D(a["week"]), wrange(a["week"]), a["jugador"], a["texto"], a["fuente"]] for a in
        sorted(F["activ"], key=lambda a: (a["week"], a["jugador"]))]
sheet("Otras anotaciones", ["Lunes", "Semana", "Jugador", "Anotación", "Fuente"], rows, [11, 22, 26, 50, 30])

# 7. Incidencias
GORD = {"alta": 0, "media": 1, "baja": 2}
rows = [[i["gravedad"], i["tipo"], D(i["semana"]) if i["semana"] else None, pretty(i["torneo"]) if i["torneo"] else "",
         i["detalle"], i["fuente"]] for i in sorted(F["incid"], key=lambda i: (GORD[i["gravedad"]], i["tipo"], i["semana"]))]
sheet("Incidencias", ["Prioridad", "Tipo", "Semana (lunes)", "Torneo", "Detalle", "Fuente"], rows,
      [9, 30, 12, 40, 90, 30])

# 8. Criterios nombres
from names import AMBIG, ROW_FIX
crit = [
    ("DIEGO", "Diego Vilches", "En la hoja de Álvaro, Diego Inostroza. Si en la misma convocatoria otra hoja pone «DIEGO V» o «DIEGO I», se usa ese."),
    ("DAVID", "Sin asignar", "Hay David Mas y David Castillo. En la hoja de Pablo se toma David Mas; en el resto, solo si otra hoja lo aclara."),
    ("ERIC", "Eric Badenes", "En la hoja de Álvaro (escuela), Eric López."),
    ("CARLA", "Carla Guerrero", "En la hoja de Álvaro, Carla Chisvert."),
    ("EUGENIA", "Eugenia Álvarez", "Salvo M15/W15 Getxo e ITF Béjar (grupo Futures, junto a Emma): Eugenia Zozaya."),
    ("MARIA", "María Ruiz", "En las listas de Álvaro/Santi. María Andrienko aparece siempre con apellido."),
    ("MIGUEL", "Miguel Uriarte", "En la hoja de Álvaro, Miguel Vidal."),
    ("JUAN", "Juan Esteban Inostroza", "Único Juan del grupo de Salva."),
    ("MARC", "Marc Martín Roca", "En convocatorias de Futures. Marc Marín (escuela) aparece siempre con apellido."),
    ("JAVI (jugadores)", "Javier Ballester", "En la columna JUGADORES. En la columna COACH, JAVI / JAVI G. = Javi Giménez."),
    ("NACHO (jugadores)", "Nacho Martínez", "En la columna COACH, NACHO = Nacho Calvo. «NACHO PARISCA» = Ignacio Parisca."),
    ("JORGE (coach)", "Jorge García / Jorge Ibáñez", "Con Victoria o en Rafa Nadal Tour → Jorge Ibáñez; en WTA, con Elina o en qualys → Jorge García."),
    ("VICTOR (coach)", "Víctor Redondo", "«VICTOR M» = Víctor M. (grupo propio)."),
    ("IA", "Ia Teporoca", ""), ("NOA", "Noa Ribera", ""), ("BONNIE", "Jinxuan Liao (Bonnie)", ""),
    ("NIK / NICK", "Xu Guilin (Nik)", ""), ("MANU", "Manuel Méndez", ""), ("ABDE", "Guennouini Abdelaziz", ""),
    ("RODRI / RODRIGO", "Rodrigo López", "También «RORIGO»."), ("VALERIA", "Valeriia Bokova", ""),
    ("VICTORIA / VIKTORIA", "Victoria Schneider", ""), ("BORRIELO", "Cristian Borriello", ""),
    ("ASMITHA / ASHMITA", "Ashmita Mitra", ""), ("AMPA GIL / AMPARO", "Amparo Gil", ""),
    ("GAO / YUANTIA GAO", "Yuantian Gao", ""), ("SANCHEZ", "Carlos Sánchez Jover", "En listas de Challenger."),
    ("ALEX / ÁLEX", "Alejandro García Carbajal", "Alex Pardo (escuela) aparece siempre con apellido."),
]
cnt = Counter()
for t in conv:
    for p in t["players"]:
        if p["certeza"].startswith("deducido"):
            cnt[p["nombre"]] += 1
sheet("Criterios nombres", ["Escrito en el Excel", "Se interpreta como", "Regla / excepciones"],
      [list(c) for c in crit], [24, 30, 100],
      note="Solo se deduce cuando el Excel pone un nombre de pila que comparten dos personas. Todo lo deducido va marcado [deducido] en las otras pestañas.")

# 9. Restos 2025
rows = [[y["hoja"], y["fila"], y["semana"], y["torneo"], y["cierre"], y["coach"], y["jugadores"]] for y in F["y25"]]
sheet("Restos 2025", ["Hoja", "Fila", "Semana", "Torneo", "Cierre", "Coach", "Jugadores"], rows,
      [12, 7, 12, 60, 16, 14, 60],
      note="Filas de 2025 que siguen en las hojas de Salva (junio-octubre 2025) y Álvaro (septiembre-diciembre 2025). No se han mezclado con 2026.")

wb.save(OUT)
print(OUT, len(conv))
