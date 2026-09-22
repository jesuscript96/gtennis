"""Excel de recogida en DOS hojas: JUGADORES y ENTRENADORES.

Todo lo que el motor usa se declara desde estas dos hojas. Las celdas ya vienen
rellenas con lo que hay hoy en la app: solo hay que corregir lo que no cuadre.
Gris = identidad (no se toca). Amarillo = se edita. Verde = fila de ejemplo.

Cada hoja lleva la ficha, la semana habitual y una semana concreta día a día.

    python manage.py plantilla_completa
    python manage.py plantilla_completa --salida /ruta/archivo.xlsx --semana 2026-09-21
"""
from datetime import date, timedelta

from django.core.management.base import BaseCommand

GRIS = "FFF2F2F2"
AMARILLO = "FFFFF6D8"
EJEMPLO = "FFEFF6EE"
CABECERA = "FF1F3A5F"
FUENTE = "Arial"
DIAS = ["LUNES", "MARTES", "MIÉRC.", "JUEVES", "VIERNES"]

# Banda de color por sección: (nombre, fill).
SEC_FICHA = ("FICHA", "FF2E5E8C")
SEC_REGLAS = ("REGLAS ESPECIALES", "FF7A4E9C")
SEC_HABITUAL = ("SEMANA HABITUAL (lo de siempre)", "FF3E7A4E")
SEC_SEMANA = ("SEMANA {} (rellenar solo lo que cambie)", "FFB5651D")

VECINDAD = {
    "CLUB": "±1 (la del club)", "SOLO": "solo su división",
    "ARRIBA": "su división y la de encima", "ABAJO": "su división y la de debajo",
    "AMBAS": "su división y las dos vecinas",
}
MOTIVOS = ["torneo", "lesión", "enfermedad", "estudios", "prueba médica", "vacaciones"]


class Command(BaseCommand):
    help = "Genera el Excel completo de recogida en dos hojas (jugadores y entrenadores)."

    def add_arguments(self, parser):
        parser.add_argument("--salida", default="docs/GTennis_datos.xlsx")
        parser.add_argument("--semana", help="Lunes de la semana a rellenar (AAAA-MM-DD).")

    def _datos(self):
        from academy.models import (Contrato, Entrenador, HorarioEntrenador,
                                    HorarioJugador, Jugador, PreferenciaPareja,
                                    PreferenciaSuperficie, Rencilla, Turno)

        turnos = list(Turno.objects.filter(activo=True).order_by("orden"))
        jugadores = list(
            Jugador.objects.filter(activo=True)
            .select_related("escuela", "division", "entrenador_responsable",
                            "turno_manana", "turno_tarde").order_by("division__nivel", "nombre"))
        entrenadores = list(
            Entrenador.objects.filter(activo=True)
            .select_related("turno_manana", "turno_tarde").order_by("nombre"))
        horarios = {(h.jugador_id, h.dia): h
                    for h in HorarioJugador.objects.select_related("turno_manana", "turno_tarde")}
        jornadas = {(h.entrenador_id, h.dia): h for h in HorarioEntrenador.objects.all()}
        rencillas = {}
        for r in Rencilla.objects.filter(activa=True).select_related("jugador_a", "jugador_b"):
            rencillas.setdefault(r.jugador_a_id, []).append(r.jugador_b.nombre)
            rencillas.setdefault(r.jugador_b_id, []).append(r.jugador_a.nombre)
        contratos = {}
        for c in Contrato.objects.filter(activo=True).select_related("entrenador"):
            contratos.setdefault(c.jugador_id, []).append(
                f"{c.entrenador.nombre}{'' if c.tipo == 'DURO' else ' (si puede)'}")
        parejas = {}
        for p in PreferenciaPareja.objects.filter(activa=True).select_related("jugador_objetivo"):
            parejas.setdefault(p.jugador_id, []).append(p.jugador_objetivo.nombre)
        superficies = {}
        for p in PreferenciaSuperficie.objects.all():
            superficies.setdefault(p.jugador_id, []).append(
                f"{p.get_superficie_display()}{' (obligatorio)' if p.estricta else ''}")
        return dict(turnos=turnos, jugadores=jugadores, entrenadores=entrenadores,
                    horarios=horarios, jornadas=jornadas, rencillas=rencillas,
                    contratos=contratos, parejas=parejas, superficies=superficies)

    def _semana_vivida(self, lunes):
        from scheduling.models import Asignacion, Disponibilidad, Semana

        s = Semana.objects.filter(fecha_inicio=lunes).first()
        if s is None:
            return {}
        datos = {}
        for a in Asignacion.objects.filter(semana=s).select_related("turno"):
            bloque = "manana" if a.turno.bloque == "MANANA" else "tarde"
            datos.setdefault(a.jugador_id, {}).setdefault(a.dia, {})[bloque] = a.turno.codigo
        for d in Disponibilidad.objects.filter(semana=s, estado="AUSENCIA_JUGADOR"):
            dia = datos.setdefault(d.jugador_id, {}).setdefault(d.dia, {})
            if d.ambito in ("DIA", "MANANA"):
                dia.setdefault("manana", "no")
            if d.ambito in ("DIA", "TARDE"):
                dia.setdefault("tarde", "no")
            if d.subtipo:
                dia["motivo"] = d.get_subtipo_display()
        return datos

    def handle(self, *args, **opts):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation

        d = self._datos()
        hoy = date.today()
        lunes = (date.fromisoformat(opts["semana"]) if opts["semana"]
                 else hoy + timedelta(days=(7 - hoy.weekday()) % 7 or 7))
        codigos = [t.codigo for t in d["turnos"]]
        man = [t.codigo for t in d["turnos"] if t.bloque == "MANANA"]
        tar = [t.codigo for t in d["turnos"] if t.bloque == "TARDE"]
        leyenda = "   ·   ".join(f"{t.codigo} = {t.hora_inicio:%H:%M}-{t.hora_fin:%H:%M}"
                                 for t in d["turnos"])

        wb = Workbook()
        wb.remove(wb.active)
        base = Font(name=FUENTE, size=11)
        blanco = Font(name=FUENTE, size=10, bold=True, color="FFFFFFFF")
        gris_txt = Font(name=FUENTE, size=10, italic=True, color="FF808080")
        titulo_f = Font(name=FUENTE, size=14, bold=True)
        borde = Border(*[Side(style="thin", color="FFD0D0D0")] * 4)
        centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
        izq = Alignment(horizontal="left", vertical="center", wrap_text=True)
        f_gris = PatternFill("solid", fgColor=GRIS)
        f_ama = PatternFill("solid", fgColor=AMARILLO)
        f_ej = PatternFill("solid", fgColor=EJEMPLO)

        def construir(titulo, intro, secciones, n_ident, filas, ejemplo, validaciones):
            """secciones: [(nombre, fill, [(col, ancho), ...]), ...]."""
            ws = wb.create_sheet(titulo)
            ws.sheet_view.showGridLines = False
            ws["A1"] = titulo
            ws["A1"].font = titulo_f
            ws["A2"] = intro
            ws["A2"].font = gris_txt
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=14)
            ws.row_dimensions[2].height = 28
            # Fila 4 = banda de sección, fila 5 = nombres de columna.
            col = 1
            for nombre, fill, cols in secciones:
                ini = col
                for etiqueta, ancho in cols:
                    ws.cell(row=5, column=col, value=etiqueta).font = blanco
                    ws.cell(row=5, column=col).fill = PatternFill("solid", fgColor=CABECERA)
                    ws.cell(row=5, column=col).alignment = centro
                    ws.cell(row=5, column=col).border = borde
                    ws.column_dimensions[get_column_letter(col)].width = ancho
                    col += 1
                banda = ws.cell(row=4, column=ini, value=nombre)
                banda.font = blanco
                banda.fill = PatternFill("solid", fgColor=fill)
                banda.alignment = centro
                if col - 1 > ini:
                    ws.merge_cells(start_row=4, start_column=ini, end_row=4, end_column=col - 1)
            fila = 6
            for i, v in enumerate(ejemplo, start=1):
                c = ws.cell(row=fila, column=i, value=v)
                c.font, c.fill, c.alignment, c.border = base, f_ej, izq, borde
            fila += 1
            for datos_fila in filas:
                for i, v in enumerate(datos_fila, start=1):
                    c = ws.cell(row=fila, column=i, value=v)
                    c.font = base
                    c.fill = f_gris if i <= n_ident else f_ama
                    c.alignment = izq if i == 1 else centro
                    c.border = borde
                fila += 1
            ws.freeze_panes = ws.cell(row=6, column=n_ident + 1)
            for letra, opciones in validaciones.items():
                dv = DataValidation(type="list", formula1='"' + ",".join(opciones) + '"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{letra}7:{letra}{fila + 200}")
            return ws

        # ================= JUGADORES =================
        ficha = [("Jugador", 26), ("Cód.", 8), ("Escuela", 16), ("División", 10),
                 ("Chico o chica", 11), ("Fecha nacim.", 13), ("Responsable", 20),
                 ("Con qué divisiones entrena", 22), ("Máx. sesiones/día", 11),
                 ("Entra el día (alta)", 13), ("Último día (baja)", 13)]
        reglas = [("NO juntar con (rencilla)", 22), ("Contrato con", 20),
                  ("Pareja fija con", 20), ("Superficie", 16), ("Pistas preferidas", 14)]
        habitual = []
        for dia in DIAS:
            habitual += [(f"{dia} mañana", 11), (f"{dia} tarde", 11)]
        semana = []
        for i, dia in enumerate(DIAS):
            fecha = lunes + timedelta(days=i)
            semana += [(f"{dia} {fecha:%d/%m} mañana", 11), ("tarde", 10), ("ausencia/torneo", 15)]
        vivida = self._semana_vivida(lunes)

        filas_j = []
        for j in d["jugadores"]:
            f = [j.nombre, j.codigo_cliente or "",
                 j.escuela.nombre if j.escuela_id else "",
                 f"D{j.division.nivel}" if j.division_id else "sin división",
                 j.get_sexo_display() if j.sexo else "",
                 f"{j.fecha_nacimiento:%d/%m/%Y}" if j.fecha_nacimiento else "",
                 j.entrenador_responsable.nombre if j.entrenador_responsable_id else "",
                 VECINDAD.get(j.vecindad, ""),
                 j.sesiones_dia_max if j.sesiones_dia_max is not None else "",
                 f"{j.fecha_alta:%d/%m/%Y}" if j.fecha_alta else "",
                 f"{j.fecha_baja:%d/%m/%Y}" if j.fecha_baja else ""]
            f += ["; ".join(d["rencillas"].get(j.id, [])), "; ".join(d["contratos"].get(j.id, [])),
                  "; ".join(d["parejas"].get(j.id, [])), "; ".join(d["superficies"].get(j.id, [])), ""]
            for i in range(5):
                h = d["horarios"].get((j.id, i))
                if h is None:
                    f += ["", ""]
                else:
                    f += [(h.turno_manana.codigo if h.turno_manana_id else "sí") if h.entrena_manana else "no",
                          (h.turno_tarde.codigo if h.turno_tarde_id else "sí") if h.entrena_tarde else "no"]
            for i in range(5):
                dia = (vivida.get(j.id) or {}).get(i, {})
                f += [dia.get("manana", ""), dia.get("tarde", ""), dia.get("motivo", "")]
            filas_j.append(f)

        ej_j = ["(ejemplo) Juan Pérez", "", "Alto Rendimiento", "D4", "Chico", "12/03/2011",
                "Mario Muniesa", "su división y la de encima", "2", "16/09/2026", "",
                "Marco Ruiz", "", "Ana Gil", "tierra (obligatorio)", "1, 2, 3",
                "M2", "T1", "M2", "no", "M2", "T1", "no", "no", "M2", "T1",
                "M2", "T1", "", "M2", "no", "torneo", "no", "no", "lesión", "M1", "no", "", "M2", "T1", ""]
        val_j = {"E": ["Chico", "Chica"], "H": list(VECINDAD.values()), "N": ["tierra", "resina"]}
        col = 17  # primera de semana habitual
        for i in range(5):
            val_j[get_column_letter(17 + 2 * i)] = man + ["no"]
            val_j[get_column_letter(18 + 2 * i)] = tar + ["no"]
        for i in range(5):
            val_j[get_column_letter(27 + 3 * i)] = man + ["no"]
            val_j[get_column_letter(28 + 3 * i)] = tar + ["no"]
            val_j[get_column_letter(29 + 3 * i)] = MOTIVOS
        construir(
            "JUGADORES",
            f"Franjas: {leyenda}.   Gris = identidad (no tocar).   Amarillo = se edita, ya trae lo que hay en la app.   "
            f"Verde = ejemplo.   Semana habitual = lo de siempre; semana concreta = solo lo que cambie esa semana.",
            [(SEC_FICHA[0], SEC_FICHA[1], ficha), (SEC_REGLAS[0], SEC_REGLAS[1], reglas),
             (SEC_HABITUAL[0], SEC_HABITUAL[1], habitual),
             (SEC_SEMANA[0].format(f"{lunes:%d/%m}"), SEC_SEMANA[1], semana)],
            2, filas_j, ej_j, val_j)

        # ================= ENTRENADORES =================
        ficha_e = [("Entrenador", 24), ("Grupo divisiones · desde", 14), ("· hasta", 10),
                   ("Franja mañana", 12), ("Franja tarde", 12),
                   ("Solo a mano (banquillo)", 13), ("Disponible esta semana", 13)]
        habitual_e = [(f"{dia} (mañana/tarde/no)", 15) for dia in DIAS]
        semana_e = []
        for i, dia in enumerate(DIAS):
            fecha = lunes + timedelta(days=i)
            semana_e += [(f"{dia} {fecha:%d/%m}", 14), ("torneo/ausencia", 15)]
        filas_e = []
        for e in d["entrenadores"]:
            f = [e.nombre, e.division_desde or "", e.division_hasta or "",
                 e.turno_manana.codigo if e.turno_manana_id else "",
                 e.turno_tarde.codigo if e.turno_tarde_id else "",
                 "sí" if e.reserva else "no", "sí" if e.disponible_semana else "no"]
            for i in range(5):
                jor = d["jornadas"].get((e.id, i))
                if jor is None:
                    f.append("")
                else:
                    partes = [b for b, v in (("mañana", jor.manana), ("tarde", jor.tarde)) if v]
                    f.append(" y ".join(partes) if partes else "no")
            for i in range(5):
                f += ["", ""]
            filas_e.append(f)
        ej_e = ["(ejemplo) Mario Muniesa", "4", "6", "M2", "", "no", "sí",
                "mañana y tarde", "mañana", "no", "mañana y tarde", "mañana",
                "sí", "", "no", "torneo", "sí", "", "sí", ""]
        val_e = {"D": man + ["cualquiera"], "E": tar + ["cualquiera"],
                 "F": ["sí", "no"], "G": ["sí", "no"]}
        for i in range(5):
            val_e[get_column_letter(8 + i)] = ["mañana y tarde", "mañana", "tarde", "no"]
            val_e[get_column_letter(13 + 2 * i)] = ["sí", "no", "mañana", "tarde"]
            val_e[get_column_letter(14 + 2 * i)] = MOTIVOS
        construir(
            "ENTRENADORES",
            f"Franjas: {leyenda}.   Grupo = divisiones que lleva (Dani Gimeno 1-2…).   "
            f"«Solo a mano» = no entra en el reparto automático pero se le puede colocar a mano.   "
            f"Semana habitual = su jornada de siempre; semana concreta = solo lo que cambie.",
            [(SEC_FICHA[0], SEC_FICHA[1], ficha_e),
             (SEC_HABITUAL[0], SEC_HABITUAL[1], habitual_e),
             (SEC_SEMANA[0].format(f"{lunes:%d/%m}"), SEC_SEMANA[1], semana_e)],
            1, filas_e, ej_e, val_e)

        wb.save(opts["salida"])
        self.stdout.write(self.style.SUCCESS(
            f"Escrito {opts['salida']} · {len(d['jugadores'])} jugadores, "
            f"{len(d['entrenadores'])} entrenadores, semana {lunes}."))
