"""Lee el Excel de recogida y lo interpreta contra la base de datos.

La premisa es que el cliente escribe como quiera —`M1+T2`, `mañana`, `8:30`,
`solo lunes`, `-`— y el trabajo de entender está aquí, no en él. Todo lo que no
se entiende no se inventa: sale en un informe con la fila, la columna y el
texto literal, para preguntar.

Escribe cuatro cosas:
  · `HorarioJugador`  — qué turnos hace cada jugador cada día (el dato que hoy
    falta y por el que el motor reparte a cuatro días en vez de a dos).
  · `Jugador.turno_manana` / `turno_tarde` / `sesiones_semana`.
  · `Entrenador.divisiones_habilitadas` y `HorarioEntrenador`.
  · `AusenciaJugador` / `VacacionesEntrenador` para lo ya conocido.

Uso:
    python manage.py importar_recogida docs/Recogida_GTennis.xlsx --dry-run
    python manage.py importar_recogida docs/Recogida_GTennis.xlsx
"""
import datetime as dt
import re
import unicodedata
from collections import Counter, defaultdict

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from academy.models import (
    Contrato, Division, Entrenador, HorarioEntrenador, HorarioJugador, Jugador,
    Pista, PreferenciaPareja, PreferenciaSuperficie, Rencilla, ResponsableJugador,
    Turno, VacacionesEntrenador,
)
from scheduling.models import Ambito, AusenciaJugador, Estado

NO_ENTRENA = {"-", "--", "no", "nada", "libre", "descansa", "descanso", "off",
              "fiesta", "ninguno", "ninguna", "0"}
VIENE_SIN_MAS = {"x", "si", "ok", "v", "vale", "viene", "*", "1"}
NUMEROS = {"un": 1, "uno": 1, "una": 1, "dos": 2, "tres": 3, "cuatro": 4,
           "cinco": 5, "seis": 6, "siete": 7, "ocho": 8, "nueve": 9, "diez": 10,
           "once": 11, "doce": 12}
MESES = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7,
         "ago": 8, "sep": 9, "set": 9, "oct": 10, "nov": 11, "dic": 12}


def numero_de(texto):
    """'4', 'cuatro', '4 sesiones', 'dos al dia' -> el primer número que salga."""
    m = re.search(r"\d+", texto)
    if m:
        return int(m.group())
    for palabra, n in NUMEROS.items():
        if re.search(rf"\b{palabra}\b", texto):
            return n
    return None


def norm(v):
    """Minúsculas, sin acentos, sin dobles espacios. None y '' dan ''."""
    if v is None:
        return ""
    s = str(v).strip()
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).lower()


class Interprete:
    """Traduce texto libre a turnos, apoyándose en el horario real del curso."""

    def __init__(self, turnos):
        self.turnos = {t.codigo.lower(): t for t in turnos}
        self.manana = [t for t in turnos if t.bloque == Turno.Bloque.MANANA]
        self.tarde = [t for t in turnos if t.bloque != Turno.Bloque.MANANA]
        # Orden dentro del bloque: "la primera de la tarde" = la más temprana.
        self.orden = {}
        for lista in (self.manana, self.tarde):
            for i, t in enumerate(sorted(lista, key=lambda x: x.hora_inicio)):
                self.orden[(t.bloque, i + 1)] = t

    def _por_hora(self, texto):
        """Turnos citados por su hora. '8:30', '8.30', '830', '15h', '15'."""
        hallados = []
        for h, m in re.findall(r"(\d{1,2})[:.h]?(\d{2})?\b", texto):
            hora = int(h)
            minuto = int(m) if m else None
            if hora > 23:
                continue
            for t in self.turnos.values():
                if t.hora_inicio.hour != hora:
                    continue
                if minuto is not None and abs(t.hora_inicio.minute - minuto) > 20:
                    continue
                hallados.append(t)
        return hallados

    def turnos_de(self, celda):
        """(turno_mañana, turno_tarde, marca). marca explica qué ha pasado."""
        s = norm(celda)
        if not s:
            return None, None, "vacio"
        if s in NO_ENTRENA:
            return None, None, "no_entrena"
        if s in VIENE_SIN_MAS:
            return None, None, "viene_sin_turno"

        man = tar = None
        # 1) Códigos explícitos: lo más fiable, manda sobre todo lo demás.
        for cod, t in self.turnos.items():
            if re.search(rf"\b{cod}\b", s):
                if t.bloque == Turno.Bloque.MANANA:
                    man = man or t
                else:
                    tar = tar or t
        # 2) Horas.
        for t in self._por_hora(s):
            if t.bloque == Turno.Bloque.MANANA:
                man = man or t
            else:
                tar = tar or t
        # 3) Ordinales atados a un bloque: "segunda de la tarde".
        ordinales = {1: r"\b(1a?|1ª|primer[ao]|primera)\b",
                     2: r"\b(2a?|2ª|segund[ao]|segunda)\b",
                     3: r"\b(3a?|3ª|tercer[ao]|tercera)\b"}
        dice_man = bool(re.search(r"\bmanana|\bmananas|\bam\b", s))
        dice_tar = bool(re.search(r"\btarde|\btardes|\bpm\b", s))
        for n, patron in ordinales.items():
            if not re.search(patron, s):
                continue
            if dice_man and man is None:
                man = self.orden.get((Turno.Bloque.MANANA, n))
            elif dice_tar and tar is None:
                tar = self.orden.get((Turno.Bloque.TARDE, n))
            elif not dice_man and not dice_tar:
                return man, tar, "ordinal_sin_bloque"
        # 4) Solo el bloque, sin turno: se resuelve luego con el resto de días.
        marca = "ok"
        if dice_man and man is None:
            marca = "manana_sin_turno"
        if dice_tar and tar is None:
            marca = "tarde_sin_turno" if marca == "ok" else "ambos_sin_turno"
        if man is None and tar is None and marca == "ok":
            return None, None, "no_entendido"
        return man, tar, marca

    def jornada(self, celda):
        """Para entrenadores: (mañana, tarde) o None si no dice nada."""
        s = norm(celda)
        if not s:
            return None
        if s in NO_ENTRENA:
            return (False, False)
        if s in VIENE_SIN_MAS or "ambas" in s or "ambos" in s or "completa" in s:
            return (True, True)
        man = bool(re.search(r"\bmanana|\bam\b|\bm[12]\b", s))
        tar = bool(re.search(r"\btarde|\bpm\b|\bt[12]\b", s))
        if not man and not tar:
            return None
        return (man, tar)


def divisiones_de(celda):
    """'1,2,3' · '1-4' · '1 a 4' · 'todas' · '2 y 3'  ->  {1,2,3,4}"""
    s = norm(celda)
    if not s or s.startswith("—") or "sin asignar" in s:
        return None
    if "todas" in s or "todo" in s:
        return set(Division.objects.values_list("nivel", flat=True))
    niveles = set()
    for a, b in re.findall(r"(\d+)\s*(?:-|a|hasta)\s*(\d+)", s):
        niveles.update(range(int(a), int(b) + 1))
    if not niveles:
        niveles = {int(n) for n in re.findall(r"\d+", s)}
    return niveles or None


def fecha_de(valor, anio_ref):
    """Acepta la fecha de Excel, '12/10/2026', '12-10-26' y '12 de octubre'."""
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    s = norm(valor)
    if not s:
        return None
    m = re.search(r"\b(\d{1,2})[/-](\d{1,2})(?:[/-](\d{2,4}))?\b", s)
    if m:
        d, mes, a = int(m.group(1)), int(m.group(2)), m.group(3)
        anio = anio_ref if a is None else (int(a) + 2000 if len(a) == 2 else int(a))
        try:
            return dt.date(anio, mes, d)
        except ValueError:
            return None
    m = re.search(r"\b(\d{1,2})\s*(?:de\s+)?([a-z]{3,})", s)
    if m and m.group(2)[:3] in MESES:
        try:
            return dt.date(anio_ref, MESES[m.group(2)[:3]], int(m.group(1)))
        except ValueError:
            return None
    return None


def columnas_de(ws, fila_cabecera=4):
    """{cabecera normalizada: índice 0-based} de la fila de cabeceras.

    Las columnas se buscan por su NOMBRE y no por su posición: la plantilla
    crece cada vez que la dirección quiere declarar algo más, y con índices
    fijos cualquier columna nueva desplazaba a las siguientes y el importador
    leía la de al lado sin enterarse.
    """
    idx = {}
    for c in ws[fila_cabecera]:
        if c.value:
            idx.setdefault(norm(str(c.value)), c.column - 1)
    return idx


def es_ejemplo(fila):
    """La plantilla trae una fila de ejemplo; no se importa."""
    v = fila[0].value
    return isinstance(v, str) and v.strip().lower().startswith("(ejemplo)")


class Command(BaseCommand):
    help = "Importa el Excel de recogida (texto libre) y explica lo que no entiende."

    def add_arguments(self, parser):
        parser.add_argument("archivo")
        parser.add_argument("--dry-run", action="store_true",
                            help="Interpreta y da el informe, sin tocar la base.")
        parser.add_argument("--anio", type=int, default=dt.date.today().year,
                            help="Año que se asume en las fechas sin año.")

    # ------------------------------------------------------------------ #
    def handle(self, *args, **opts):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise CommandError("Falta openpyxl (pip install openpyxl).")

        wb = load_workbook(opts["archivo"], data_only=True)
        self.dudas = []          # (hoja, fila, columna, texto, motivo)
        self.anio = opts["anio"]
        turnos = list(Turno.objects.filter(activo=True).order_by("orden"))
        if not turnos:
            raise CommandError("No hay turnos activos: revisa el curso cargado.")
        self.interprete = Interprete(turnos)
        self._jugadores_idx = {norm(j.nombre): j
                               for j in Jugador.objects.filter(activo=True)}
        self._entrenadores_idx = {norm(e.nombre): e
                                  for e in Entrenador.objects.filter(activo=True)}

        with transaction.atomic():
            r_j = self._jugadores(wb)
            r_e = self._entrenadores(wb)
            r_a = self._ausencias(wb)
            if opts["dry_run"]:
                transaction.set_rollback(True)

        self._informe(r_j, r_e, r_a, opts["dry_run"])

    def _duda(self, hoja, fila, col, texto, motivo):
        self.dudas.append((hoja, fila, col, str(texto), motivo))

    # ------------------------------------------------------------------ #
    def _jugadores(self, wb):
        if "Jugadores" not in wb.sheetnames:
            return {}
        ws = wb["Jugadores"]
        res = Counter()
        nombres = self._jugadores_idx
        col = columnas_de(ws)
        dias_col = [col.get(norm(d)) for d in ("LUNES", "MARTES", "MIÉRCOLES",
                                               "JUEVES", "VIERNES")]

        for fila in ws.iter_rows(min_row=5, values_only=False):
            nombre = fila[0].value
            if not nombre or es_ejemplo(fila):
                continue
            jug = nombres.get(norm(nombre))
            if jug is None:
                self._duda("Jugadores", fila[0].row, "A", nombre, "jugador no encontrado")
                res["sin_jugador"] += 1
                continue

            # 1) Leer los cinco días.
            lecturas = {}
            for d, indice in enumerate(dias_col):
                if indice is None or indice >= len(fila):
                    lecturas[d] = (None, None, "vacio")
                    continue
                celda = fila[indice]
                man, tar, marca = self.interprete.turnos_de(celda.value)
                lecturas[d] = (man, tar, marca)
                if marca == "no_entendido":
                    self._duda("Jugadores", celda.row, celda.column_letter,
                               celda.value, "no sé qué turno es")
                elif marca == "ordinal_sin_bloque":
                    self._duda("Jugadores", celda.row, celda.column_letter,
                               celda.value, "¿primera de la mañana o de la tarde?")

            # 2) Resolver los «mañana» sueltos con lo que hace el resto de días.
            habitual_m = Counter(m.id for m, _, _ in lecturas.values() if m)
            habitual_t = Counter(t.id for _, t, _ in lecturas.values() if t)
            def preferido(cont, bloque):
                if cont:
                    return Turno.objects.get(pk=cont.most_common(1)[0][0])
                candidatos = (self.interprete.manana if bloque == Turno.Bloque.MANANA
                              else self.interprete.tarde)
                return sorted(candidatos, key=lambda t: t.hora_inicio)[0]

            for d, (man, tar, marca) in lecturas.items():
                if marca in ("manana_sin_turno", "ambos_sin_turno") and man is None:
                    man = preferido(habitual_m, Turno.Bloque.MANANA)
                    res["deducido"] += 1
                if marca in ("tarde_sin_turno", "ambos_sin_turno") and tar is None:
                    tar = preferido(habitual_t, Turno.Bloque.TARDE)
                    res["deducido"] += 1
                lecturas[d] = (man, tar, marca)

            # 3) Escribir. Solo los días con respuesta: sin fila, el motor elige.
            for d, (man, tar, marca) in lecturas.items():
                if marca in ("vacio", "viene_sin_turno", "no_entendido",
                             "ordinal_sin_bloque"):
                    HorarioJugador.objects.filter(jugador=jug, dia=d).delete()
                    continue
                HorarioJugador.objects.update_or_create(
                    jugador=jug, dia=d,
                    # Lo que la plantilla dice sin franja ya viene deducido
                    # arriba, así que aquí un bloque vacío es «no entrena».
                    defaults={"turno_manana": man, "turno_tarde": tar,
                              "entrena_manana": man is not None,
                              "entrena_tarde": tar is not None},
                )
                res["dias"] += 1
                if marca == "no_entrena":
                    res["dias_libres"] += 1

            # 4) Turno habitual en la ficha, para los días sin declarar.
            if habitual_m:
                jug.turno_manana = preferido(habitual_m, Turno.Bloque.MANANA)
            if habitual_t:
                jug.turno_tarde = preferido(habitual_t, Turno.Bloque.TARDE)

            # 5) Sesiones por semana + control contra los días declarados.
            celda_ses = self._celda(fila, col, "Ses./semana")
            ses = norm(celda_ses.value)
            declaradas = sum(bool(m) + bool(t) for m, t, _ in lecturas.values())
            if ses:
                num = numero_de(ses)
                if num is not None:
                    jug.sesiones_semana = num
                    if declaradas and abs(num - declaradas) > 1:
                        self._duda("Jugadores", celda_ses.row, celda_ses.column_letter, celda_ses.value,
                                   f"dice {num} pero los días suman {declaradas}")
                else:
                    self._duda("Jugadores", celda_ses.row, celda_ses.column_letter, celda_ses.value,
                               "no es un número")
            # 6) Lo que se declara una vez y vale para toda la semana: la
            #    franja de siempre pisa a lo deducido de los días, y las fechas
            #    de alta y baja marcan desde y hasta cuándo cuenta.
            for etiqueta, campo, bloque in (
                ("De normal · MAÑANA", "turno_manana", Turno.Bloque.MANANA),
                ("De normal · TARDE", "turno_tarde", Turno.Bloque.TARDE),
            ):
                celda = self._celda(fila, col, etiqueta)
                texto = norm(celda.value) if celda is not None else ""
                if not texto:
                    continue
                man, tar, marca = self.interprete.turnos_de(celda.value)
                elegido = man if bloque == Turno.Bloque.MANANA else tar
                if elegido is not None:
                    setattr(jug, campo, elegido)
                elif marca == "no_entrena":
                    setattr(jug, campo, None)
                elif marca not in ("vacio", "viene"):
                    self._duda("Jugadores", celda.row, celda.column_letter,
                               celda.value, "no sé qué franja es")

            campos = ["turno_manana", "turno_tarde", "sesiones_semana"]
            celda = self._celda(fila, col, "Máx. al día")
            if celda is not None and norm(celda.value):
                n = numero_de(norm(celda.value))
                if n:
                    jug.sesiones_dia_max = n
                    campos.append("sesiones_dia_max")
                else:
                    self._duda("Jugadores", celda.row, celda.column_letter,
                               celda.value, "no es un número")
            for etiqueta, campo in (("Entra el día", "fecha_alta"),
                                    ("Último día", "fecha_baja")):
                celda = self._celda(fila, col, etiqueta)
                if celda is None or not norm(celda.value):
                    continue
                f = fecha_de(celda.value, self.anio)
                if f is None:
                    self._duda("Jugadores", celda.row, celda.column_letter,
                               celda.value, "no es una fecha")
                else:
                    setattr(jug, campo, f)
                    campos.append(campo)
                    res[campo] += 1

            jug.save(update_fields=campos)
            res["jugadores"] += 1

            self._extras(fila, col, jug, res)
        return res

    # ------------------------------------------------------------------ #
    @staticmethod
    def _celda(fila, col, nombre):
        """La celda de esa columna en esta fila, o None si la plantilla no la
        trae (una versión vieja del fichero, o la han borrado)."""
        i = col.get(norm(nombre))
        return fila[i] if i is not None and i < len(fila) else None

    def _extras(self, fila, col, jug, res):
        """Las columnas de la derecha: solo si aplica, casi siempre vacías."""
        def c(nombre):
            return self._celda(fila, col, nombre)

        # Escuela corregida.
        celda = c("Escuela correcta")
        if celda is not None and norm(celda.value):
            from academy.models import Escuela

            esc = next((e for e in Escuela.objects.all()
                        if norm(e.nombre) == norm(celda.value)), None)
            if esc is None:
                self._duda("Jugadores", celda.row, celda.column_letter,
                           celda.value, "no reconozco esa escuela")
            elif esc != jug.escuela:
                jug.escuela = esc
                jug.save(update_fields=["escuela"])
                res["escuela"] += 1

        # División corregida.
        celda_div = c("División correcta")
        div = norm(celda_div.value) if celda_div is not None else ""
        if div:
            n = numero_de(div)
            d = Division.objects.filter(nivel=n).first() if n else None
            if d is None:
                self._duda("Jugadores", celda_div.row, celda_div.column_letter,
                           celda_div.value, "no reconozco esa división")
            elif d != jug.division:
                jug.division = d
                jug.save(update_fields=["division"])
                res["division"] += 1

        # Responsable corregido.
        celda_resp = c("Responsable correcto")
        nombre = norm(celda_resp.value) if celda_resp is not None else ""
        if nombre and not nombre.startswith("—"):
            ent = self._entrenador(nombre)
            if ent is None:
                self._duda("Jugadores", celda_resp.row, celda_resp.column_letter,
                           celda_resp.value, "ese entrenador no existe")
            else:
                ResponsableJugador.objects.filter(jugador=jug).update(activo=False)
                ResponsableJugador.objects.update_or_create(
                    jugador=jug, entrenador=ent,
                    defaults={"activo": True, "prioridad": 1},
                )
                # El responsable de la ficha es el mismo dato: si no, la app
                # diría una cosa y el motor otra.
                jug.entrenador_responsable = ent
                jug.save(update_fields=["entrenador_responsable"])
                res["responsable"] += 1

        # Pareja preferida. «a poder ser» / «si puede» lo baja a preferente.
        celda_par = c("Entrena SIEMPRE con")
        for nom in self._nombres(celda_par.value if celda_par is not None else None):
            otro = self._jugador(nom)
            if otro is None or otro.id == jug.id:
                self._duda("Jugadores", celda_par.row, celda_par.column_letter, nom,
                           "no encuentro a ese jugador")
                continue
            blando = re.search(r"a poder ser|si puede|preferib|mejor si",
                               norm(celda_par.value))
            PreferenciaPareja.objects.update_or_create(
                jugador=jug, jugador_objetivo=otro,
                defaults={"tipo": PreferenciaPareja.Tipo.SOFT if blando
                          else PreferenciaPareja.Tipo.HARD, "activa": True},
            )
            res["parejas"] += 1

        # Veto.
        celda_veto = c("NO ponerlo con")
        celda_nota = c("Notas")
        for nom in self._nombres(celda_veto.value if celda_veto is not None else None):
            otro = self._jugador(nom)
            if otro is None or otro.id == jug.id:
                self._duda("Jugadores", celda_veto.row, celda_veto.column_letter, nom,
                           "no encuentro a ese jugador")
                continue
            a, b = sorted([jug, otro], key=lambda x: x.id)
            Rencilla.objects.update_or_create(
                jugador_a=a, jugador_b=b,
                defaults={"activa": True, "motivo": str(
                    (celda_nota.value if celda_nota is not None else "") or "")[:200]},
            )
            res["vetos"] += 1

        # Contrato de patrocinio.
        celda_con = c("Contrato con")
        for nom in self._nombres(celda_con.value if celda_con is not None else None):
            ent = self._entrenador(norm(nom))
            if ent is None:
                self._duda("Jugadores", celda_con.row, celda_con.column_letter, nom,
                           "ese entrenador no existe")
                continue
            Contrato.objects.update_or_create(
                jugador=jug, entrenador=ent, defaults={"activo": True})
            res["contratos"] += 1

        # Superficie.
        celda_sup = c("Superficie")
        sup = norm(celda_sup.value) if celda_sup is not None else ""
        if sup:
            if "tierra" in sup or "batida" in sup:
                elegida = Pista.Superficie.TIERRA
            elif "resina" in sup or "rapid" in sup or "dura" in sup:
                elegida = Pista.Superficie.RESINA
            else:
                elegida = None
                self._duda("Jugadores", celda_sup.row, celda_sup.column_letter,
                           celda_sup.value, "¿tierra o resina?")
            if elegida:
                PreferenciaSuperficie.objects.update_or_create(
                    jugador=jug, superficie=elegida,
                    defaults={"estricta": "estricta" in sup or "solo" in sup})
                res["superficie"] += 1

        nota = celda_nota.value if celda_nota is not None else None
        if nota and str(nota).strip():
            self._duda("Jugadores", celda_nota.row, celda_nota.column_letter, nota,
                       "nota para leer a mano")

    # ------------------------------------------------------------------ #
    def _nombres(self, celda):
        """'Pepe y Juan' / 'Pepe, Juan' -> ['pepe', 'juan']. Vacío -> []."""
        s = norm(celda)
        if not s or s in NO_ENTRENA:
            return []
        s = re.sub(r"\b(a poder ser|si puede|preferiblemente|mejor si|con)\b", " ", s)
        return [t.strip() for t in re.split(r",|;|\by\b|\+|/", s) if t.strip()]

    def _jugador(self, nombre):
        return self._buscar(nombre, self._jugadores_idx)

    def _entrenador(self, nombre):
        return self._buscar(nombre, self._entrenadores_idx)

    @staticmethod
    def _buscar(nombre, indice):
        """Exacto primero; si no, el único que contenga todas las palabras."""
        clave = norm(nombre)
        if clave in indice:
            return indice[clave]
        palabras = clave.split()
        if not palabras:
            return None
        candidatos = [v for k, v in indice.items()
                      if all(p in k.split() for p in palabras)]
        return candidatos[0] if len(candidatos) == 1 else None

    # ------------------------------------------------------------------ #
    def _entrenadores(self, wb):
        if "Entrenadores" not in wb.sheetnames:
            return {}
        ws = wb["Entrenadores"]
        res = Counter()
        nombres = self._entrenadores_idx
        col = columnas_de(ws)
        dias_col = [col.get(norm(d)) for d in ("LUNES", "MARTES", "MIÉRCOLES",
                                               "JUEVES", "VIERNES")]

        for fila in ws.iter_rows(min_row=5, values_only=False):
            nombre = fila[0].value
            if not nombre or es_ejemplo(fila):
                continue
            ent = nombres.get(norm(nombre))
            if ent is None:
                self._duda("Entrenadores", fila[0].row, "A", nombre, "entrenador no encontrado")
                continue

            celda_div = self._celda(fila, col, "Divisiones correctas")
            niveles = divisiones_de(celda_div.value) if celda_div is not None else set()
            if niveles:
                divs = list(Division.objects.filter(nivel__in=niveles))
                faltan = niveles - {d.nivel for d in divs}
                if faltan:
                    self._duda("Entrenadores", celda_div.row, celda_div.column_letter,
                               celda_div.value,
                               f"divisiones que no existen: {sorted(faltan)}")
                ent.divisiones_habilitadas.set(divs)
                res["divisiones"] += 1

            # Su franja fija, igual que la del alumno: en blanco entra donde
            # haga falta.
            campos = []
            for etiqueta, campo, bloque in (
                ("De normal · MAÑANA", "turno_manana", Turno.Bloque.MANANA),
                ("De normal · TARDE", "turno_tarde", Turno.Bloque.TARDE),
            ):
                celda = self._celda(fila, col, etiqueta)
                if celda is None or not norm(celda.value):
                    continue
                man, tar, marca = self.interprete.turnos_de(celda.value)
                elegido = man if bloque == Turno.Bloque.MANANA else tar
                if elegido is not None:
                    setattr(ent, campo, elegido)
                    campos.append(campo)
                elif marca == "no_entrena":
                    setattr(ent, campo, None)
                    campos.append(campo)
                else:
                    self._duda("Entrenadores", celda.row, celda.column_letter,
                               celda.value, "no sé qué franja es")
            if campos:
                ent.save(update_fields=campos)
                res["franjas"] += 1

            for d, indice in enumerate(dias_col):
                if indice is None or indice >= len(fila):
                    continue
                celda = fila[indice]
                jor = self.interprete.jornada(celda.value)
                if jor is None:
                    HorarioEntrenador.objects.filter(entrenador=ent, dia=d).delete()
                    continue
                HorarioEntrenador.objects.update_or_create(
                    entrenador=ent, dia=d,
                    defaults={"manana": jor[0], "tarde": jor[1]},
                )
                res["jornada"] += 1
            celda_nota = self._celda(fila, col, "Notas")
            nota = celda_nota.value if celda_nota is not None else None
            if nota and str(nota).strip():
                self._duda("Entrenadores", celda_nota.row, celda_nota.column_letter,
                           nota, "nota para leer a mano")
            res["entrenadores"] += 1
        return res

    # ------------------------------------------------------------------ #
    def _ausencias(self, wb):
        if "Ausencias" not in wb.sheetnames:
            return {}
        ws = wb["Ausencias"]
        res = Counter()
        jug, ent = self._jugadores_idx, self._entrenadores_idx
        col = columnas_de(ws)

        for fila in ws.iter_rows(min_row=5, values_only=False):
            quien = fila[0].value
            if not quien or es_ejemplo(fila):
                continue
            c_desde = self._celda(fila, col, "Desde")
            c_hasta = self._celda(fila, col, "Hasta")
            desde = fecha_de(c_desde.value if c_desde is not None else None, self.anio)
            hasta = fecha_de(c_hasta.value if c_hasta is not None else None,
                             self.anio) or desde
            if desde is None:
                self._duda("Ausencias", fila[0].row, "B",
                           c_desde.value if c_desde is not None else "", "no es una fecha")
                continue
            if hasta < desde:
                self._duda("Ausencias", c_hasta.row, c_hasta.column_letter,
                           c_hasta.value, "el «hasta» va antes que el «desde»")
                continue
            c_motivo = self._celda(fila, col, "Motivo")
            motivo = str((c_motivo.value if c_motivo is not None else "") or "").strip()
            c_nota = self._celda(fila, col, "Notas")
            nota = str((c_nota.value if c_nota is not None else "") or "").strip()
            if nota:
                motivo = f"{motivo} · {nota}" if motivo else nota
            clave = norm(quien)

            if clave in jug:
                # «Qué se pierde» es la columna nueva; se acepta también el
                # nombre viejo para los ficheros ya repartidos.
                ambito = self._ambito(
                    self._celda(fila, col, "Qué se pierde")
                    or self._celda(fila, col, "Turnos afectados")
                )
                AusenciaJugador.objects.update_or_create(
                    jugador=jug[clave], fecha_inicio=desde, fecha_fin=hasta,
                    ambito=ambito,
                    defaults={"estado": Estado.AUSENCIA_JUGADOR, "nota": motivo},
                )
                res["jugadores"] += 1
            elif clave in ent:
                VacacionesEntrenador.objects.update_or_create(
                    entrenador=ent[clave], fecha_inicio=desde, fecha_fin=hasta,
                    defaults={"motivo": motivo},
                )
                res["entrenadores"] += 1
            else:
                self._duda("Ausencias", fila[0].row, "A", quien, "no está en jugadores ni entrenadores")
        return res

    def _ambito(self, celda):
        if celda is None:
            return Ambito.DIA
        s = norm(celda.value)
        if not s:
            return Ambito.DIA
        for cod in ("m1", "m2", "jp", "t1", "t2"):
            if re.search(rf"\b{cod}\b", s):
                return getattr(Ambito, cod.upper())
        if "manana" in s:
            return Ambito.MANANA
        if "tarde" in s:
            return Ambito.TARDE
        self._duda("Ausencias", celda.row, celda.column_letter, celda.value,
                   "no sé qué se pierde")
        return Ambito.DIA

    # ------------------------------------------------------------------ #
    def _informe(self, r_j, r_e, r_a, dry):
        w = self.stdout.write
        w(self.style.MIGRATE_HEADING(
            "\nLo que he entendido" + ("  (ENSAYO: nada guardado)" if dry else "")))
        w(f"  Jugadores leídos ........ {r_j.get('jugadores', 0)}")
        w(f"    días con turno ........ {r_j.get('dias', 0)}"
          f"   (de los cuales {r_j.get('dias_libres', 0)} son «no viene»)")
        w(f"    turnos deducidos ...... {r_j.get('deducido', 0)}"
          "   (decía solo «mañana»/«tarde»)")
        w(f"    escuelas corregidas ... {r_j.get('escuela', 0)}")
        w(f"    divisiones corregidas . {r_j.get('division', 0)}")
        w(f"    responsables fijados .. {r_j.get('responsable', 0)}")
        w(f"    altas / bajas ......... {r_j.get('fecha_alta', 0)}"
          f" / {r_j.get('fecha_baja', 0)}")
        w(f"    parejas / vetos ....... {r_j.get('parejas', 0)} / {r_j.get('vetos', 0)}")
        w(f"    contratos ............. {r_j.get('contratos', 0)}")
        w(f"    superficie ............ {r_j.get('superficie', 0)}")
        w(f"  Entrenadores leídos ..... {r_e.get('entrenadores', 0)}")
        w(f"    divisiones fijadas .... {r_e.get('divisiones', 0)}")
        w(f"    franjas fijadas ....... {r_e.get('franjas', 0)}")
        w(f"    días de jornada ....... {r_e.get('jornada', 0)}")
        w(f"  Ausencias de jugador .... {r_a.get('jugadores', 0)}")
        w(f"  Ausencias de entrenador . {r_a.get('entrenadores', 0)}")

        if not self.dudas:
            w(self.style.SUCCESS("\nSin dudas: se entendió todo.\n"))
            return
        por_motivo = defaultdict(list)
        for hoja, f, c, texto, motivo in self.dudas:
            por_motivo[motivo].append((hoja, f, c, texto))
        w(self.style.WARNING(f"\nPara preguntar  ({len(self.dudas)} celdas)"))
        for motivo, casos in sorted(por_motivo.items(), key=lambda x: -len(x[1])):
            w(f"\n  {motivo}  ({len(casos)})")
            for hoja, f, c, texto in casos[:12]:
                w(f"    {hoja}!{c}{f}   «{texto[:60]}»")
            if len(casos) > 12:
                w(f"    … y {len(casos) - 12} más")
        w("")
